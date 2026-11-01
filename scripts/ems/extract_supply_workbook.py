"""Reads the Captain 5 supply room workbook and writes the tracked extraction used by `php artisan ems:apply-supply-workbook`.

    "C:\\...\\python.exe" -I scripts\\ems\\extract_supply_workbook.py --workbook "C:\\...\\MEDICAL_SUPPLY_LIST_with_Count_Tab.xlsx"

Inputs: sheet "Supply Room Count" (starting physical counts) and "Supply Room Inventory" (shelves, order levels,
package sizes). Other sheets are listed but never read for data. The workbook is opened for reading and never saved.

Outputs (repeatable, no timestamps):
  database/data/ems/supply_room_workbook_2026.json   schema mbfd.ems.supply_room_workbook.v1
  docs/ems/supply-room-workbook-exceptions.md        the report a fire captain reads

Rows are matched to the Hub's supply items by exact normalized text only (whitespace collapsed, nothing else) against the
tracked captain5_supply_room_inventory.json. There is no similarity score anywhere. A row that matches nothing must have a
reviewed decision in supply_workbook_reviewed.py or the run stops.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import sys
from pathlib import Path

# -I removes the script folder from sys.path; add our own sibling modules back explicitly.
sys.path.insert(0, str(Path(__file__).resolve().parent))

import openpyxl  # noqa: E402

from supply_workbook_parsing import (  # noqa: E402
    PACKAGE_UNITS, cell_text, norm, parse_order, parse_package, parse_quantity, parse_threshold, plural, threshold_view,
)
from supply_workbook_report import write_report  # noqa: E402
from supply_workbook_reviewed import REVIEWED_OVERRIDES, REVIEWED_UNMATCHED  # noqa: E402

SCHEMA = "mbfd.ems.supply_room_workbook.v1"
COUNT_SHEET = "Supply Room Count"
INVENTORY_SHEET = "Supply Room Inventory"
REPO = Path(__file__).resolve().parents[2]
DEFAULT_TRACKED = REPO / "database" / "data" / "ems" / "captain5_supply_room_inventory.json"
DEFAULT_JSON = REPO / "database" / "data" / "ems" / "supply_room_workbook_2026.json"
DEFAULT_DOC = REPO / "docs" / "ems" / "supply-room-workbook-exceptions.md"


# --------------------------------------------------------------------------- reading the workbook

def read_sheet(ws, last_col: int, first_row: int) -> list[dict]:
    """Every non-empty row as {row, cells: [A..], } with raw cell values."""
    rows = []
    for r in range(first_row, ws.max_row + 1):
        cells = [ws.cell(r, c).value for c in range(1, last_col + 1)]
        if any(v is not None and str(v).strip() != "" for v in cells):
            rows.append({"row": r, "cells": cells})
    return rows


def classify(rows: list[dict], item_cols: int, merged: set[int]) -> tuple[list[dict], list[dict], int]:
    """Split sheet rows into item rows and headings. Returns (items, headings, loc_only_count).

    Heading: exactly one filled cell, and it is either text in column B, or text in column A that is not a shelf
    code used on any item row. A row with only a shelf code is a placeholder, not an item and not a heading.
    """
    shelf_codes = {norm(r["cells"][0]) for r in rows if any(v is not None for v in r["cells"][1:item_cols]) and r["cells"][0] is not None}
    items, headings, loc_only = [], [], 0
    block = None
    for r in rows:
        filled = [i for i, v in enumerate(r["cells"][:item_cols]) if v is not None and str(v).strip() != ""]
        if len(filled) == 1 and isinstance(r["cells"][filled[0]], str):
            idx, text = filled[0], norm(r["cells"][filled[0]])
            if idx == 1 or (idx == 0 and text not in shelf_codes):
                block = text
                headings.append({"row": r["row"], "cell": "AB"[idx], "text": text, "merged_band": r["row"] in merged})
                continue
            if idx == 0:
                loc_only += 1
                continue
        r["block"] = block
        items.append(r)
    return items, headings, loc_only


def merged_rows(ws) -> set[int]:
    out: set[int] = set()
    for rng in ws.merged_cells.ranges:
        out.update(range(rng.min_row, rng.max_row + 1))
    return out


def no_formulas(path: str, sheets: tuple[str, ...]) -> None:
    wb = openpyxl.load_workbook(path, data_only=False)
    for name in sheets:
        for row in wb[name].iter_rows(max_col=10):
            for cell in row:
                if cell.data_type == "f":
                    raise SystemExit(f"{name}!{cell.coordinate} holds a formula; this import reads typed values only.")


# --------------------------------------------------------------------------- matching

class Index:
    """Exact normalized text -> Hub source_row_key, in three fixed tiers (cell text, name + alias, name)."""

    def __init__(self, rows: list[dict]):
        self.tiers: list[dict[str, list[str]]] = [{}, {}, {}]
        for r in rows:
            name, alias = norm(r["item_text"]), norm(r.get("bag_alias_text"))
            for tier, text in zip(self.tiers, (norm(r["item_cell_text"]), norm(f"{name} {alias}") if alias else "", name)):
                if text:
                    tier.setdefault(text, []).append(r["row_key"])

    def find(self, text: str) -> dict:
        for method, tier in zip(("cell_text", "name_and_alias", "name"), self.tiers):
            hits = tier.get(text, [])
            if len(hits) == 1:
                return {"source_row_key": hits[0], "method": method}
            if len(hits) > 1:
                return {"source_row_key": None, "method": "ambiguous", "candidates": sorted(hits)}
        return {"source_row_key": None, "method": "none"}


def apply_override(sheet: str, row: int, text: str) -> dict | None:
    for o in REVIEWED_OVERRIDES:
        if o["sheet"] == sheet and o["row"] == row and o["text"] == text:
            return {"source_row_key": o["source_row_key"], "method": "reviewed_override", "reason": o["reason"]}
    return None


# --------------------------------------------------------------------------- building the extraction

def stored_as(qty: int) -> str:
    return f"Order level stored as {qty} individual unit{'' if qty == 1 else 's'}."


def exception(group: str, sheet: str | None, row: int | None, key: str | None, text: str | None, detail: str, action: str) -> dict:
    return {"group": group, "sheet": sheet, "row": row, "source_row_key": key, "item_text": text, "detail": detail, "action": action}


def build(workbook: str, tracked_path: Path) -> dict:
    sha = hashlib.sha256(Path(workbook).read_bytes()).hexdigest()
    no_formulas(workbook, (COUNT_SHEET, INVENTORY_SHEET))
    wb = openpyxl.load_workbook(workbook, data_only=True)
    tracked_raw = tracked_path.read_bytes()
    tracked = json.loads(tracked_raw)
    tracked_by_key = {r["row_key"]: r for r in tracked["rows"]}
    index = Index(tracked["rows"])

    ws_c, ws_i = wb[COUNT_SHEET], wb[INVENTORY_SHEET]
    if [norm(ws_c.cell(1, c).value) for c in (1, 2, 3)] != ["LOC", "ITEM", "QUANTITY"]:
        raise SystemExit("Count sheet header is not LOC / ITEM / QUANTITY on row 1.")
    inv_header = [norm(ws_i.cell(2, c).value) for c in (1, 2, 3, 5, 6, 7, 8, 9, 10)]
    if inv_header != ["LOC", "ITEM", "ORDER LEVEL", "QTY PER BOX/CS", "PRIMARY", "SECONDARY", "TERTIARY", "ITEM #", "ORDER AMOUNT"]:
        raise SystemExit(f"Inventory sheet header row 2 changed: {inv_header}")
    stray = [ (c.coordinate) for row in ws_i.iter_rows(min_row=1, min_col=11) for c in row if c.value is not None]
    stray += [c.coordinate for row in ws_c.iter_rows(min_row=1, min_col=4) for c in row if c.value is not None]
    if stray:
        raise SystemExit(f"Unexpected values outside the known columns: {stray[:10]}")

    count_items, count_heads, count_loc_only = classify(read_sheet(ws_c, 3, 3), 3, merged_rows(ws_c))
    inv_items, inv_heads, inv_loc_only = classify(read_sheet(ws_i, 10, 3), 10, merged_rows(ws_i))

    exceptions: list[dict] = []
    reviewed_hits: set[tuple[str, int]] = set()

    # ---- Count rows
    count_rows = []
    for r in count_items:
        a, b, c = r["cells"][:3]
        text = norm(b) if b is not None else None
        raw = int(c) if isinstance(c, float) and c.is_integer() else c
        row = {"row": r["row"], "block": r.get("block"), "loc": cell_text(a), "item_text": text, "quantity_raw": raw}
        row["parsed"] = parse_quantity(c)
        if text is None:
            row["match"] = {"source_row_key": None, "method": "none"}
        else:
            row["match"] = apply_override("count", r["row"], text) or index.find(text)
        count_rows.append(row)

    # ---- Inventory rows
    inv_rows = []
    for r in inv_items:
        a, b, c, d, e, f, g, h, i, j = r["cells"][:10]
        text = norm(b) if isinstance(b, str) else None
        row = {
            "row": r["row"], "block": r.get("block"), "loc": cell_text(a), "item_text": text if text else cell_text(b),
            "order_level_text": cell_text(c), "check_text": cell_text(d), "qty_per_package_text": cell_text(e),
            "item_number_text": cell_text(i), "order_amount_text": cell_text(j),
        }
        row["parsed"] = {"threshold": parse_threshold(row["order_level_text"]), "package": parse_package(e), "order": parse_order(j)}
        row["match"] = (apply_override("inventory", r["row"], text) or index.find(text)) if text else {"source_row_key": None, "method": "none"}
        inv_rows.append(row)

    # ---- duplicates: a Hub item may be named by one row per sheet
    for sheet, rows in (("count", count_rows), ("inventory", inv_rows)):
        seen: dict[str, dict] = {}
        for row in rows:
            key = row["match"]["source_row_key"]
            if key is None:
                continue
            if key in seen:
                for bad in (seen[key], row):
                    bad["match"] = {"source_row_key": None, "method": "duplicate", "candidates": [key]}
                exceptions.append(exception("duplicate_match", sheet, row["row"], key, row["item_text"],
                                            f"Rows {seen[key]['row']} and {row['row']} on the {sheet} sheet name the same item.", "Neither row is used for this item."))
            else:
                seen[key] = row

    # ---- unmatched rows need a reviewed decision
    inv_unmatched_by_text = {r["item_text"]: r for r in inv_rows if r["match"]["source_row_key"] is None and r["item_text"]}
    new_items: list[dict] = []
    for sheet, rows in (("count", count_rows), ("inventory", inv_rows)):
        for row in rows:
            if row["match"]["source_row_key"] is not None or row["match"]["method"] == "duplicate":
                continue
            decision = REVIEWED_UNMATCHED.get((sheet, row["row"]))
            if decision is None:
                raise SystemExit(f"{sheet} row {row['row']} ({row['item_text']!r}) matches no Hub item and has no reviewed decision in supply_workbook_reviewed.py.")
            reviewed_hits.add((sheet, row["row"]))
            row["disposition"] = decision["action"]
            if decision["action"] == "new":
                partner = inv_unmatched_by_text.get(row["item_text"]) if sheet == "count" else None
                row["match"] = {"source_row_key": f"wb.c{row['row']}", "method": "new_item"}
                if partner is not None:
                    partner["match"] = {"source_row_key": f"wb.c{row['row']}", "method": "new_item"}
                    partner["disposition"] = "new"
                new_items.append({"count_row": row, "inventory_row": partner, "decision": decision})
            else:
                exceptions.append(exception("unmatched_row", sheet, row["row"], None, row["item_text"] or str(row.get("quantity_raw")),
                                            decision["reason"], "No item created; no count applied."))
    unused = set(REVIEWED_UNMATCHED) - reviewed_hits
    if unused:
        raise SystemExit(f"Reviewed decisions no longer match any row (the workbook changed): {sorted(unused)}")

    # ---- per item plan
    by_key_count = {r["match"]["source_row_key"]: r for r in count_rows if r["match"]["source_row_key"]}
    by_key_inv = {r["match"]["source_row_key"]: r for r in inv_rows if r["match"]["source_row_key"]}
    plans: list[dict] = []
    changes: list[dict] = []
    for tr in tracked["rows"]:
        plans.append(plan_item(tr["row_key"], tr, by_key_count.get(tr["row_key"]), by_key_inv.get(tr["row_key"]), exceptions, changes))
    for n in new_items:
        plans.append(plan_new(n, exceptions))

    # ---- blank quantities on the Count sheet and quantity rows of headings
    for row in count_rows:
        if row["item_text"] and row["parsed"]["kind"] == "blank":
            exceptions.append(exception("blank_quantity", "count", row["row"], row["match"]["source_row_key"], row["item_text"],
                                        "The item is on the Count sheet with no quantity.", "Stays Not counted (never zero) until the next count."))

    exceptions.sort(key=lambda e: (e["group"], e["sheet"] or "", e["row"] or 0))
    populated = [r for r in count_rows if r["item_text"] and r["parsed"]["kind"] != "blank"]
    summary = {
        "count_sheet_item_rows": len([r for r in count_rows if r["item_text"]]),
        "count_sheet_populated_rows": len(populated),
        "count_sheet_blank_quantity_rows": len([r for r in count_rows if r["item_text"] and r["parsed"]["kind"] == "blank"]),
        "count_sheet_quantity_without_item": len([r for r in count_rows if not r["item_text"]]),
        "count_sheet_headings": len(count_heads),
        "count_sheet_shelf_only_rows": count_loc_only,
        "inventory_sheet_item_rows": len(inv_rows),
        "inventory_sheet_headings": len(inv_heads),
        "hub_items_matched_on_count": len([k for k in by_key_count if k in tracked_by_key]),
        "hub_items_matched_on_inventory": len([k for k in by_key_inv if k in tracked_by_key]),
        "new_items": len(new_items),
        "baselines": len([p for p in plans if p["count"] and p["count"]["baseline_base"] is not None]),
        "exceptions_by_group": {g: len([e for e in exceptions if e["group"] == g]) for g in sorted({e["group"] for e in exceptions})},
    }
    sheets = [{"name": ws.title, "dimensions": ws.dimensions, "role": {COUNT_SHEET: "count", INVENTORY_SHEET: "inventory"}.get(ws.title, "not_an_input")} for ws in wb.worksheets]
    return {
        "schema": SCHEMA,
        "source": {
            "file_name": Path(workbook).name, "sha256": sha, "sheets": sheets,
            "tracked_supply_source_sha256": tracked["source"]["sha256"],
            "tracked_supply_extraction_sha256": hashlib.sha256(tracked_raw).hexdigest(),
            "row_counts": summary,
            "notes": [
                "Cell text is stored with runs of whitespace collapsed to one space, as in the tracked supply room extraction.",
                "Supplier names and prices (PRIMARY, SECONDARY, TERTIARY) are not copied.",
                "Matching is exact normalized text against the tracked supply room extraction; there is no similarity matching.",
                "ORDER AMOUNT is recorded for reference only and is never used as an order level.",
                "Cell comments are not imported (one exists, on Inventory B125, about how three medications are packaged together).",
            ],
        },
        "headings": [{"sheet": "count", **h} for h in count_heads] + [{"sheet": "inventory", **h} for h in inv_heads],
        "count_rows": count_rows,
        "inventory_rows": inv_rows,
        "items": plans,
        "exceptions": exceptions,
        "_changes": changes,
    }


def plan_item(key: str, tr: dict, count: dict | None, inv: dict | None, exceptions: list[dict], changes: list[dict]) -> dict:
    """What the workbook says about one existing Hub item."""
    controlled = bool(tr["flags"].get("controlled_substance"))
    db_pkg = tr["parsed"]["package"]
    wb_pkg = inv["parsed"]["package"] if inv else {"unit": None, "base_units_per_package": None}
    unit = wb_pkg["unit"] if wb_pkg["unit"] in ("each", *PACKAGE_UNITS) else None
    unit = unit or (db_pkg["unit"] if db_pkg["unit"] in ("each", *PACKAGE_UNITS) else None)
    factor = wb_pkg["base_units_per_package"] or db_pkg["base_units_per_package"]
    name = norm(tr["item_cell_text"])

    thr = threshold_view(inv["order_level_text"] if inv else tr["order_level_text"], unit, factor, controlled)
    thr_db = db_threshold(tr)
    shelf = plan_shelf(tr, count, inv)

    item = {
        "source_row_key": key, "new": False,
        "count_row": count["row"] if count else None, "inventory_row": inv["row"] if inv else None,
        "item_text": name, "shelf": shelf,
        "threshold": thr,
        "package": {"unit": wb_pkg["unit"], "units_per_package": wb_pkg["base_units_per_package"], "text": inv["qty_per_package_text"] if inv else None},
        "count": None, "count_unit": "each",
    }
    if count and count["item_text"]:
        item["count"] = resolve_count(count["parsed"], thr, unit, factor, controlled)
    item["count_unit"] = choose_count_unit(item["count"], thr, unit, factor)
    if item["count"] is not None and item["count"].get("count_unit") is None:
        item["count"]["count_unit"] = item["count_unit"]

    # exception records
    where = ("count", count["row"]) if count else (None, None)
    c = item["count"]
    if c and c["exception"]:
        group = {"ambiguous_bare_quantity": "ambiguous_bare_quantity", "controlled_no_baseline": "controlled_no_baseline"}.get(c["exception"], "unit_conversion")
        exceptions.append(exception(group, *where, key, name, c["message"], c["action"]))
    if thr["status"] == "assumed_each":
        exceptions.append(exception("threshold_no_unit", "inventory", inv["row"] if inv else None, key, name, thr["reason"], stored_as(thr["qty_base"])))
    elif thr["status"] == "not_converted":
        exceptions.append(exception("threshold_not_converted", "inventory", inv["row"] if inv else None, key, name, thr["reason"], "No order level is stored for this item until someone enters one."))
    elif thr["status"] == "as_needed":
        exceptions.append(exception("threshold_as_needed", "inventory", inv["row"] if inv else None, key, name, thr["reason"], "No order level is stored."))
    elif thr["status"] == "explicit_equivalent":
        exceptions.append(exception("threshold_note", "inventory", inv["row"] if inv else None, key, name, thr["reason"], stored_as(thr["qty_base"])))
    if shelf["conflict"]:
        exceptions.append(exception("shelf_conflict", "inventory", inv["row"] if inv else None, key, name,
                                    f"Inventory sheet shelf {shelf['inventory_loc']}, Count sheet shelf {shelf['count_loc']}.", f"Shelf stays {shelf['code']} (the Inventory sheet)."))
    if shelf["source"] == "count":
        exceptions.append(exception("shelf_from_count", "inventory", inv["row"] if inv else None, key, name,
                                    f"The Inventory sheet has no shelf; the Count sheet says {shelf['count_loc']}.", f"Shelf set to {shelf['count_loc']}."))
    # differences from what the Hub holds today (report only)
    diffs = []
    if shelf["code"] and shelf["code"] != (tr["location_code"] or None) and shelf["source"] == "inventory":
        diffs.append(("Shelf", tr["location_code"], shelf["code"]))
    if not controlled and thr["qty_base"] != thr_db:
        diffs.append(("Order level (individual units)", thr_db, thr["qty_base"]))
    if inv and norm(inv["order_level_text"]) != norm(tr["order_level_text"]):
        diffs.append(("Order level text", tr["order_level_text"], inv["order_level_text"]))
    if inv and norm(inv["qty_per_package_text"]) != norm(tr["qty_per_package_text"]):
        diffs.append(("Quantity per package text", tr["qty_per_package_text"], inv["qty_per_package_text"]))
    for field, before, after in diffs:
        changes.append({"source_row_key": key, "item_text": name, "field": field, "before": before, "after": after})
    return item


def db_threshold(tr: dict) -> int | None:
    """The threshold the Hub holds today, by the SupplyRowMapper rule on the tracked row."""
    t, p = tr["parsed"]["threshold"], tr["parsed"]["package"]
    if not t or t.get("qty") is None or t.get("unit") is None:
        return None
    if t.get("comparison") not in (None, "lt"):
        return None
    if t["unit"] == "each":
        return t["qty"]
    if t["unit"] == p.get("unit") and p.get("base_units_per_package"):
        return t["qty"] * p["base_units_per_package"]
    return None


def plan_shelf(tr: dict, count: dict | None, inv: dict | None) -> dict:
    inv_loc = inv["loc"] if inv else None
    cnt_loc = count["loc"] if count else None
    conflict = bool(inv_loc and cnt_loc and inv_loc != cnt_loc)
    if inv_loc:
        return {"code": inv_loc, "source": "inventory", "inventory_loc": inv_loc, "count_loc": cnt_loc, "conflict": conflict}
    if cnt_loc:
        return {"code": cnt_loc, "source": "count", "inventory_loc": None, "count_loc": cnt_loc, "conflict": False}
    return {"code": None, "source": None, "inventory_loc": None, "count_loc": None, "conflict": False}


def resolve_count(parsed: dict, thr: dict, package_unit: str | None, factor: int | None, controlled: bool) -> dict:
    """The rules for turning a Count sheet cell into a starting number. Mirrored in SupplyWorkbookCounts.php."""
    out = {"kind": parsed["kind"], "value": parsed["value"], "unit": parsed["unit"], "units_note": parsed["units_note"],
           "count_unit": None, "baseline_base": None, "exception": None, "message": None, "action": None}
    if controlled:
        return {**out, "exception": "controlled_no_baseline", "message": "Controlled substance with a count on the sheet.", "action": "Controlled items never get a starting count."}
    kind, value, unit = parsed["kind"], parsed["value"], parsed["unit"]
    if kind == "blank":
        return out
    if kind == "invalid":
        return {**out, "exception": "unreadable_quantity", "message": "The quantity cell cannot be read as a number or as a number with a box, case, pack or bag.", "action": "No starting count applied."}
    if kind == "bare":
        if thr["class"] == "package":
            usable = thr["unit"] if thr["unit"] == package_unit and factor else None
            return {**out, "count_unit": usable or "each", "exception": "ambiguous_bare_quantity",
                    "message": f"The sheet shows {value}, but the order level ({thr['text']}) is written in {plural(thr['unit'])}, so {value} could be {plural(thr['unit'])} or single units.",
                    "action": "No starting count applied; Monday's count sets it." + (f" Count unit set to {usable}." if usable else "")}
        return {**out, "count_unit": "each", "baseline_base": value}
    # explicit unit
    if unit == "each":
        return {**out, "count_unit": "each", "baseline_base": value}
    if unit == package_unit and factor:
        if parsed["units_note"] is not None and parsed["units_note"] != factor:
            return {**out, "exception": "unit_note_mismatch", "message": f"The sheet says {value} {unit} ({parsed['units_note']}), but the item is {factor} per {unit}.", "action": "No starting count applied."}
        return {**out, "count_unit": unit, "baseline_base": value * factor}
    if package_unit is None:
        why = "the item has no package size on the Inventory sheet"
    elif package_unit == "each":
        why = "the Inventory sheet lists this item as single units, not boxes or cases"
    elif package_unit != unit:
        why = f"the Inventory sheet packs this item by the {package_unit}, not the {unit}"
    else:
        why = f"the number of units per {unit} is not stated"
    return {**out, "exception": "unit_mismatch", "message": f"The sheet says {value} {unit}, but {why}.", "action": "No starting count applied."}


def choose_count_unit(count: dict | None, thr: dict, package_unit: str | None, factor: int | None) -> str:
    """What the person types on Monday: each, or the item's package unit when that is how it is counted and ordered."""
    if count and count.get("count_unit"):
        return count["count_unit"]
    if thr["class"] == "package" and thr["unit"] == package_unit and factor:
        return thr["unit"]
    return "each"


def plan_new(n: dict, exceptions: list[dict]) -> dict:
    """A new supply item from a Count row (and the matching Inventory row when it has one)."""
    row, inv, decision = n["count_row"], n["inventory_row"], n["decision"]
    spec = decision["item"]
    key = f"wb.c{row['row']}"
    name = spec["name"]
    wb_pkg = inv["parsed"]["package"] if inv else {"unit": None, "base_units_per_package": None}
    unit = wb_pkg["unit"] if wb_pkg["unit"] in ("each", *PACKAGE_UNITS) else None
    factor = wb_pkg["base_units_per_package"]
    thr = threshold_view(inv["order_level_text"] if inv else None, unit, factor, False)
    cnt_loc = row["loc"]
    inv_loc = inv["loc"] if inv else None
    shelf = {"code": inv_loc or cnt_loc, "source": "inventory" if inv_loc else ("count" if cnt_loc else None), "inventory_loc": inv_loc, "count_loc": cnt_loc, "conflict": bool(inv_loc and cnt_loc and inv_loc != cnt_loc)}
    count = resolve_count(row["parsed"], thr, unit, factor, False)
    reasons = ["Added from the supply room workbook; it is not on the printed 2021 list."]
    if spec.get("extra_review"):
        reasons.append(spec["extra_review"])
    if shelf["code"] is None:
        reasons.append("No shelf is given on the workbook.")
    if thr["status"] == "none":
        reasons.append("No order level is given on the workbook.")
    item = {
        "source_row_key": key, "new": True, "count_row": row["row"], "inventory_row": inv["row"] if inv else None,
        "item_text": row["item_text"], "shelf": shelf, "threshold": thr,
        "package": {"unit": wb_pkg["unit"], "units_per_package": wb_pkg["base_units_per_package"], "text": inv["qty_per_package_text"] if inv else None},
        "count": count, "count_unit": "each",
        "new_item": {
            "name": name, "bag_alias_text": spec.get("bag_alias_text"), "section_key": spec["section_key"], "catalog_ref": spec["catalog_ref"],
            "refrigerated": bool(spec.get("refrigerated", False)), "controlled_substance": False,
            "item_numbers": [inv["item_number_text"]] if inv and inv["item_number_text"] else [],
            "qty_per_package_text": inv["qty_per_package_text"] if inv else None,
            "order_amount_text": inv["order_amount_text"] if inv else None,
            "needs_review": True, "review_reasons": reasons, "reason": decision["reason"],
        },
    }
    item["count_unit"] = choose_count_unit(count, thr, unit, factor)
    count["count_unit"] = count["count_unit"] or item["count_unit"]
    exceptions.append(exception("new_item", "count", row["row"], key, row["item_text"], decision["reason"], "Created as a new supply item" + (" with a starting count." if count["baseline_base"] is not None else ".")))
    if count["exception"]:
        exceptions.append(exception("ambiguous_bare_quantity" if count["exception"] == "ambiguous_bare_quantity" else "unit_conversion", "count", row["row"], key, name, count["message"], count["action"]))
    if thr["status"] == "assumed_each":
        exceptions.append(exception("threshold_no_unit", "inventory", inv["row"], key, name, thr["reason"], stored_as(thr["qty_base"])))
    return item


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--workbook", required=True, help="Path to MEDICAL_SUPPLY_LIST_with_Count_Tab.xlsx")
    parser.add_argument("--tracked", default=str(DEFAULT_TRACKED))
    parser.add_argument("--out", default=str(DEFAULT_JSON))
    parser.add_argument("--doc", default=str(DEFAULT_DOC))
    args = parser.parse_args()

    result = build(args.workbook, Path(args.tracked))
    changes = result.pop("_changes")
    Path(args.out).write_text(json.dumps(result, indent=2, ensure_ascii=False) + "\n", encoding="utf-8", newline="\n")
    write_report(result, changes, Path(args.doc))
    print(json.dumps(result["source"]["row_counts"], indent=2))
    print(f"Wrote {args.out}")
    print(f"Wrote {args.doc}")


if __name__ == "__main__":
    main()
